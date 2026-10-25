from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from zipfile import ZipFile, ZIP_DEFLATED
import re,json,hashlib,shutil

root=Path(r'C:/sistema tributario dayanex')
out=root/'Documentacion Word';out.mkdir(exist_ok=True)
sources=[('README.md','01 Guia de inicio'),('backend/README.md','02 Guia del backend'),('sistema_tributario_rimpe/README.md','03 Guia de Flutter'),('backend/SRI.md','04 Integracion SRI y alcance'),('VERIFICACION.md','05 Verificacion del sistema'),('PRODUCCION.md','06 Publicacion y produccion'),('VERCEL.md','07 Despliegue en Vercel'),('reportes/qa/artifact.md','08 Registro de formato del reporte')]
manifest=[]
def plain(s):
    def link(m):
        target=(path.parent/m[2]).resolve()
        for rel,name in sources:
            if target==(root/rel).resolve(): return name+'.docx'
        return m[1]+' ('+m[2]+')'
    s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',link,s)
    for old,new in [('backend/SRI.md','04 Integracion SRI y alcance.docx'),('SRI.md','04 Integracion SRI y alcance.docx'),('VERCEL.md','07 Despliegue en Vercel.docx'),('VERIFICACION.md','05 Verificacion del sistema.docx')]:s=s.replace(old,new)
    return s.replace('**','').replace('`','')
def shade(cell,color):
    el=OxmlElement('w:shd');el.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(el)
for rel,name in sources:
    path=root/rel; source=path.read_text(encoding='utf-8-sig'); d=Document()
    sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85)
    for sn in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        st=d.styles[sn];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(5)
    for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.styles['Title'].font.size=Pt(20)
    d.add_paragraph(name[3:],'Title')
    if rel=='sistema_tributario_rimpe/README.md':
        d.add_paragraph('Actualización del 24 de octubre de 2026: la aplicación web ya se compiló y probó. Las sesiones web se mantienen solo en memoria y las retenciones se obtienen de comprobantes importados. Las notas anteriores se conservan a continuación como historial; no ejecutar migrate dev ni flutter create para iniciar el proyecto existente. Para el arranque actual, consultar 01 Guia de inicio; para publicar, 07 Despliegue en Vercel.')
    elif rel=='PRODUCCION.md':
        d.add_paragraph('Para el destino elegido, consultar 07 Despliegue en Vercel. La configuración de servidor propio se conserva como alternativa y requiere adaptar el proxy HTTPS.')
    lines=source.splitlines();i=0;code=False;covered=[]
    while i<len(lines):
        line=lines[i];i+=1
        if line.startswith('```'):code=not code;continue
        if code:
            p=d.add_paragraph();p.paragraph_format.space_after=Pt(0)
            r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(8)
            covered.append(line);continue
        if not line.strip():continue
        if not line.startswith(('#','|')):
            while i<len(lines) and lines[i].strip() and not re.match(r'^(#{1,6}\s|\||```|- |\d+\. )',lines[i]):
                line+=' '+lines[i].strip();i+=1
        if line.startswith('|'):
            block=[line]
            while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
            records=[[plain(c.strip()) for c in row.strip().strip('|').split('|')] for row in block if not re.match(r'^\|[\s:|\-]+\|$',row)]
            table=d.add_table(rows=0,cols=max(map(len,records)));table.style='Table Grid'
            borders=OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
            table._tbl.tblPr.append(borders)
            for ri,record in enumerate(records):
                cells=table.add_row().cells
                for ci,s in enumerate(record):
                    cells[ci].text=s;covered.append(s)
                    shade(cells[ci],'D9E5F2' if ri==0 else ('F4F6F8' if ri%2==0 else 'FFFFFF'))
                    for p in cells[ci].paragraphs:
                        for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
                if ri==0:
                    h=OxmlElement('w:tblHeader');table.rows[ri]._tr.get_or_add_trPr().append(h)
            d.add_paragraph();continue
        m=re.match(r'^(#{1,6})\s+(.*)',line)
        if m:
            s=plain(m[2]);d.add_paragraph(s,'Heading '+str(min(len(m[1]),3)));covered.append(s)
        else:
            s=plain(line)
            if s.startswith('- '):s=s[2:];d.add_paragraph(s,'List Bullet')
            else:d.add_paragraph(s)
            covered.append(s)
    dest=out/(name+'.docx');d.save(dest)
    check=Document(dest)
    body='\n'.join(p.text for p in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for r in t.rows for c in r.cells)
    assert all(s in body for s in covered if s),rel
    manifest.append({'source':rel,'output':str(dest),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'content_verified':True})
shutil.copy2(root/'reportes/Reporte semanal segunda entrega 19 al 23 octubre 2026.docx',out/'09 Reporte semanal segunda entrega.docx')
with ZipFile(out/'Respaldo documental original.zip','w',ZIP_DEFLATED) as z:
    for rel,_ in sources:z.write(root/rel,rel)
(root/'reportes/qa/conversion.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(manifest,ensure_ascii=False))
