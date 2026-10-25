from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out=Path(r'C:\sistema tributario dayanex\Documentacion Word\10 Guia paso a paso Vercel y base de datos.docx')
d=Document(); s=d.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.7); s.left_margin=s.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Title'].font.size=Pt(24)
for b in list(d.styles.element.iter(qn('w:pBdr'))): b.getparent().remove(b)
s.header.paragraphs[0].text='DAYANEX   |   Guía de publicación'
s.header.paragraphs[0].style='Caption'
f=s.footer.paragraphs[0]; f.text='Guía preparada el 25 de octubre de 2026  •  Página '
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,1)
def code(t):
 for line in t.splitlines():
  a=d.add_paragraph(); a.paragraph_format.space_after=Pt(2)
  r=a.add_run(line); r.font.name='Consolas'; r.font.size=Pt(9)
def page(t): d.add_page_break(); h(t)
def table(rows,widths):
 t=d.add_table(rows=0, cols=len(widths)); t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for c,w,txt in zip(cells,widths,row):
   c.width=Inches(w); c.text=txt
   pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   pr.append(borders)
   if i==0:
    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EEF5'); pr.append(shade)
   for a in c.paragraphs:
    a.paragraph_format.space_before=Pt(5); a.paragraph_format.space_after=Pt(5)
    for r in a.runs: r.font.size=Pt(10); r.bold=(i==0)
  if i==0:
   repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)

d.add_paragraph('Publicar Dayanex en Vercel', 'Title')
p('Guía paso a paso para crear la base de datos y poner en línea la aplicación desde Windows.')
p('Este documento corresponde al proyecto ubicado en C:\\sistema tributario dayanex. La ruta propuesta usa Vercel para la página y la API, y Neon para PostgreSQL. No necesitas tener un repositorio en GitHub para seguir estos pasos.')
h('1 Qué debes abrir')
table([['Lugar','Para qué sirve','Dirección o acceso'],['Vercel','Publicar y revisar errores de la aplicación','https://vercel.com'],['Neon','Crear y consultar PostgreSQL','https://console.neon.tech'],['PowerShell','Preparar archivos y ejecutar comandos','Inicio de Windows → PowerShell']], [1,2.5,3.4])
p('El navegador se conecta a Vercel por HTTPS. La API alojada en Vercel se conecta a Neon mediante TLS. La base local de Docker no es accesible desde Vercel y no se trasladará automáticamente.')
h('2 Antes de comenzar')
p('Crea una cuenta en Vercel y otra en Neon con tu correo. Activa la verificación en dos pasos donde esté disponible. Revisa los límites, costos y condiciones del plan antes de contratarlo; esta guía no presupone un plan gratuito apto para cualquier uso.')
p('Abre PowerShell y ejecuta cada línea por separado. Si aparece un error, detente y resuélvelo antes de continuar. No copies las explicaciones como comandos.')
code('node --version\nnpm.cmd --version\nflutter.bat --version')
p('Usa Node.js 22 para coincidir con la configuración propuesta. Si un comando no se reconoce, instala Node.js desde nodejs.org o configura Flutter siguiendo docs.flutter.dev, y vuelve a abrir PowerShell.')
p('Alcance actual: la publicación en la nube todavía requiere validación. El módulo SRI continúa en ambiente de pruebas; subir la aplicación no lo habilita para emitir comprobantes reales. No existe una garantía de seguridad del 100 %.')

page('3 Crear la base de datos en Neon')
p('1. Entra a https://console.neon.tech e inicia sesión. Crea un proyecto nuevo con un nombre reconocible, por ejemplo dayanex. Elige una región cercana a la que utilizarás para las funciones de Vercel.')
p('2. Abre el proyecto y selecciona su rama principal. Conserva la base y el usuario que crea Neon si no necesitas otros. Anota sus nombres, sin guardar contraseñas en este documento.')
p('3. Pulsa Connect. Selecciona la rama, base de datos y usuario correctos. Copia la cadena PostgreSQL, que comienza con postgresql://, no el comando psql completo.')
p('4. Obtén dos cadenas para la misma base: una con Connection pooling desactivado, para crear las tablas; otra con pooling activado, para la aplicación. El host de esta última normalmente incluye -pooler.')
p('5. En ambas cadenas conserva los parámetros existentes y verifica sslmode=require. Agrega &sslaccept=strict si aún no aparece. Si la cadena no tiene parámetros, el primero se introduce con ? y los siguientes con &. No dupliques sslmode ni sslaccept.')
code('.../neondb?sslmode=require&sslaccept=strict')
p('Este fragmento es ilustrativo; no sirve como conexión real. Mantén el usuario, contraseña, host y base que entrega Neon. Guarda las dos cadenas en un gestor de contraseñas. No las compartas por chat, correo ni capturas.')
h('Dónde ver la base y sus tablas')
p('Dentro del proyecto Neon abre SQL Editor. Selecciona la misma rama y base usadas en Connect. Después de ejecutar las migraciones del paso 5, pega esta consulta y pulsa Run:')
code("SELECT table_name\nFROM information_schema.tables\nWHERE table_schema = 'public'\nORDER BY table_name;")
p('Antes de las migraciones puede no haber tablas de la aplicación. Después deben aparecer tablas como users, companies y roles. Para contar registros sin mostrar datos personales puedes ejecutar:')
code('SELECT COUNT(*) AS usuarios FROM users;\nSELECT COUNT(*) AS empresas FROM companies;')
p('No hace falta instalar pgAdmin para consultar la base. SQL Editor es suficiente. Las migraciones crean la estructura; no copian los usuarios, empresas o comprobantes que tienes en tu computadora.')

page('4 Preparar la aplicación en tu computadora')
p('En PowerShell prepara el backend. Los comandos siguientes se ejecutan dentro de la carpeta del proyecto existente:')
code("Set-Location 'C:\\sistema tributario dayanex\\backend'\nnpm.cmd ci\nnpm.cmd run prisma:generate\nnpm.cmd run build")
p('Compila la interfaz Flutter y cópiala a la carpeta public del backend. El valor vacío de API_BASE_URL permite que la interfaz use la API del mismo dominio.')
code("Set-Location 'C:\\sistema tributario dayanex\\sistema_tributario_rimpe'\nflutter.bat pub get\nflutter.bat build web --release --dart-define=API_BASE_URL=\nNew-Item -ItemType Directory -Path '..\\backend\\public' -Force\nCopy-Item -Path '.\\build\\web\\*' `\n  -Destination '..\\backend\\public' -Recurse -Force\nTest-Path '..\\backend\\public\\index.html'")
p('La última línea debe mostrar True. Copia las dos líneas de Copy-Item juntas; el acento grave al final de la primera indica continuación en PowerShell.')
h('5 Crear las tablas y los permisos')
p('Introduce la conexión DIRECTA de Neon cuando se solicite. La entrada se oculta. Solo se establece en esta sesión de PowerShell; no reemplaza el archivo .env local.')
code("Set-Location 'C:\\sistema tributario dayanex\\backend'\n$dbSecret = Read-Host 'Conexion DIRECTA de Neon' -AsSecureString\n$env:DATABASE_URL = [System.Net.NetworkCredential]::new(\n  '', $dbSecret).Password\nnpm.cmd run db:deploy\nnpm.cmd run seed\nRemove-Item Env:DATABASE_URL\nRemove-Variable dbSecret")
p('db:deploy debe terminar sin errores y seed debe completar la carga de roles y permisos. seed no crea tu cuenta de acceso. No uses migrate reset ni db push para sustituir este procedimiento en una base con información.')
p('Si algún comando falla, elimina la variable DATABASE_URL con la línea indicada y cierra la terminal al terminar. Abre SQL Editor en Neon y verifica las tablas con la consulta de la página anterior.')

page('6 Crear el proyecto de Vercel')
p('Desde backend ejecuta los siguientes comandos. npx puede pedir permiso para descargar la herramienta Vercel; acepta si el paquete indicado es vercel. El inicio de sesión abrirá el navegador.')
code("Set-Location 'C:\\sistema tributario dayanex\\backend'\nnpx.cmd vercel@latest login\nnpx.cmd vercel@latest link")
p('En las preguntas de link selecciona tu cuenta, elige crear un proyecto nuevo si no existe, escribe un nombre como dayanex y usa ./ como directorio. Estás vinculando backend, no la carpeta completa que contiene los reportes Word.')
p('En el panel de Vercel abre ese proyecto y entra en Settings. Comprueba la detección de NestJS, Node.js 22 y el comando de compilación de vercel.json. No establezcas dist como Output Directory de un sitio estático. Esta aplicación incluye una API.')
p('Abre Settings → Domains y anota el dominio estable asignado al proyecto. Si aún no hay uno, configura un dominio disponible antes de completar CORS_ORIGIN. No supongas que el nombre del proyecto garantiza un dominio específico.')
h('7 Agregar las variables')
p('En Settings → Environment Variables añade las siguientes variables para Production. Pega los valores sin comillas exteriores. Los cambios requieren un nuevo despliegue.')
table([['Variable','Valor que debes ingresar'],['NODE_ENV','production'],['DATABASE_URL','Conexión de Neon CON pooling, sslmode=require y sslaccept=strict.'],['JWT_SECRET','Secreto aleatorio generado con el comando inferior.'],['CORS_ORIGIN','Origen real: https://tu-dominio.vercel.app, sin / final ni rutas.'],['ALLOW_PUBLIC_REGISTRATION','false al comenzar.']], [2.15,4.75])
p('Genera JWT_SECRET en tu computadora y guárdalo de forma privada. El comando produce 96 caracteres hexadecimales:')
code('node -e "console.log(require(\'crypto\').randomBytes(48).toString(\'hex\'))"')
p('No agregues estas claves al código ni a Flutter. No necesitas configurar PORT ni VERCEL manualmente. Si más adelante usas Preview, configura un entorno y una base separados; no conectes pruebas a datos reales por comodidad.')

page('8 Publicar y crear tu primer acceso')
p('Antes de habilitar el registro, entra en Settings → Deployment Protection. Comprueba que la protección elegida cubra también el dominio de producción; la disponibilidad depende del plan. Una protección solo de previews no basta.')
p('Si tu plan no permite proteger producción, mantén ALLOW_PUBLIC_REGISTRATION=false y solicita implementar un alta administrativa controlada antes de crear la primera cuenta en público. No abras el registro para todo Internet como sustituto.')
p('Con las variables guardadas, publica desde backend:')
code("Set-Location 'C:\\sistema tributario dayanex\\backend'\nnpx.cmd vercel@latest --prod")
p('En Vercel → Deployments abre el despliegue. Espera el estado Ready. Si falla, abre Build Logs; si compila pero la API falla, revisa los registros de ejecución en Logs. Comparte únicamente el error, ocultando claves y conexiones.')
p('Abre tu dominio por HTTPS y luego /api/health. Tras pasar la protección de Vercel, la API debe responder con status igual a ok y database igual a connected. Comprueba también que cargue la pantalla de inicio de sesión.')
h('Crear la cuenta bajo protección')
p('1. Comprueba en una ventana privada, sin sesión de Vercel, que el despliegue está bloqueado para visitantes. Solo entonces cambia ALLOW_PUBLIC_REGISTRATION a true y vuelve a ejecutar el despliegue.')
p('2. Accede tú con la autorización de Vercel, abre el registro de la aplicación y crea tu cuenta con una contraseña única de al menos 12 caracteres. La cuenta de Vercel y la de Dayanex son distintas.')
p('3. Cambia ALLOW_PUBLIC_REGISTRATION a false y despliega nuevamente. Verifica que un registro nuevo sea rechazado con 403 y que puedas iniciar sesión con la cuenta creada.')
p('4. Conserva la protección mientras validas las funciones. Revisa y retira los despliegues anteriores que mantengan el registro habilitado antes de quitar la protección: cambiar una variable no modifica versiones ya publicadas.')
h('HTTPS y estado de la aplicación')
p('Vercel administra HTTPS para sus dominios. Si añades un dominio propio, configura sus DNS según Settings → Domains y espera que quede verificado; actualiza CORS_ORIGIN y despliega otra vez.')
p('NODE_ENV=production activa controles del servidor, pero no cambia el ambiente del SRI. Antes de usar datos tributarios reales faltan la validación funcional en la nube, la revisión de dependencias y la validación de los procesos fiscales pendientes.')

page('9 Resolver errores y mantener la publicación')
table([['Qué ocurre','Qué revisar'],['Error de PostgreSQL o TLS','DATABASE_URL de Production debe ser la de Neon, con sslmode=require y sslaccept=strict. Nunca localhost.'],['Faltan tablas o permisos','Ejecuta db:deploy y seed con la conexión directa a la misma rama y base.'],['Error de JWT_SECRET','Usa el secreto aleatorio real de al menos 48 caracteres y vuelve a desplegar.'],['Página vacía o no encontrada','Verifica backend/public/index.html, repite la compilación y copia de Flutter, y publica desde backend.'],['Error CORS','Usa el origen HTTPS exacto desde el que accedes, sin barra final. Actualiza y despliega.'],['Registro responde 403','Es lo esperado con registro cerrado. Sigue el procedimiento protegido del paso 8.'],['Vercel solicita acceso','Puede ser Deployment Protection, independiente del inicio de sesión de Dayanex.']], [2.15,4.75])
h('Actualizar y respaldar')
p('Para publicar cambios, repite la compilación del backend y de Flutter, la copia a public y vercel --prod. Si hay migraciones nuevas, respalda primero la base y revisa su impacto antes de ejecutar db:deploy. Un rollback del código no revierte cambios de la base.')
p('En Neon revisa las opciones de restauración y su retención según tu plan. Define copias periódicas y prueba una restauración en una base separada. Revisa el cifrado en reposo y los controles de acceso ofrecidos por el proveedor; HTTPS por sí solo no cifra todos los datos almacenados.')
p('Mantén las cuentas con segundo factor, limita quién accede a Neon y Vercel, y revisa errores y consumo. Hay dependencias y procesos SRI pendientes de revisión; no presentes esta publicación como certificada o invulnerable.')
h('Fuentes y comprobación final')
p('Documentación oficial consultada el 25 de octubre de 2026. Los nombres de los menús pueden variar:')
for url in ['https://neon.com/docs/get-started/signing-up','https://neon.com/docs/connect/connect-from-any-app','https://vercel.com/docs/frameworks/backend/nestjs','https://vercel.com/docs/environment-variables','https://vercel.com/docs/deployment-protection','https://vercel.com/docs/cdn-security/encryption']:
 a=d.add_paragraph(url); a.paragraph_format.space_after=Pt(3)
 for r in a.runs:r.font.size=Pt(9)
p('Comprobación final: despliegue Ready, interfaz visible, health conectado, acceso verificado, registro cerrado, secretos privados y restauración prevista. Esta guía prepara el proceso; no confirma que tus cuentas o servicios ya estén creados.')
d.save(out)
print(out)
